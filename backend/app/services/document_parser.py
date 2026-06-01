"""
文档解析服务
将上传的课程资料文件 (PDF/DOCX/PPTX/MD/TXT) 解析为纯文本
"""

import os
from pathlib import Path
from loguru import logger


# 支持的文件类型及对应扩展名
SUPPORTED_TYPES = {
    "pdf": ".pdf",
    "docx": ".docx",
    "pptx": ".pptx",
    "md": ".md",
    "txt": ".txt",
}


def detect_file_type(filename: str) -> str:
    """
    根据文件扩展名检测文件类型
    :param filename: 原始文件名
    :return: 文件类型标识 (pdf/docx/pptx/md/txt)
    :raises ValueError: 不支持的文件类型
    """
    ext = Path(filename).suffix.lower()
    for file_type, supported_ext in SUPPORTED_TYPES.items():
        if ext == supported_ext:
            return file_type
    raise ValueError(f"不支持的文件类型: {ext}，支持的类型: {list(SUPPORTED_TYPES.values())}")


def parse_pdf(file_path: str) -> str:
    """
    解析 PDF 文件，提取文本内容
    使用 PyMuPDF (fitz) 逐页提取, 保留分页信息作为分隔符
    :param file_path: PDF 文件路径
    :return: 提取后的纯文本
    """
    import fitz  # PyMuPDF
    doc = fitz.open(file_path)
    pages_text: list[str] = []
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text()
        if text.strip():
            # 在每页文本前添加页码标记，便于后续切片时追踪来源
            pages_text.append(f"[第{page_num}页]\n{text.strip()}")
    doc.close()
    result = "\n\n".join(pages_text)
    logger.info(f"PDF 解析完成: {file_path}, 共 {len(pages_text)} 页文本")
    return result


def parse_docx(file_path: str) -> str:
    """
    解析 DOCX 文件，提取文本内容
    按段落提取, 过滤空段落
    :param file_path: DOCX 文件路径
    :return: 提取后的纯文本
    """
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # 保留标题样式信息作为标记
            if para.style.name.startswith("Heading"):
                level = para.style.name.split()[-1]
                paragraphs.append(f"{'#' * int(level)} {text}")
            else:
                paragraphs.append(text)

    # 也提取表格中的文本
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    result = "\n\n".join(paragraphs)
    logger.info(f"DOCX 解析完成: {file_path}, 共 {len(paragraphs)} 段文本")
    return result


def parse_pptx(file_path: str) -> str:
    """
    解析 PPTX 文件，提取文本内容
    逐页幻灯片提取, 包含标题和正文
    :param file_path: PPTX 文件路径
    :return: 提取后的纯文本
    """
    from pptx import Presentation
    prs = Presentation(file_path)
    slides_text: list[str] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_parts: list[str] = [f"[幻灯片 {slide_num}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_parts.append(text)
        if len(slide_parts) > 1:  # 有实际文本内容
            slides_text.append("\n".join(slide_parts))
    result = "\n\n".join(slides_text)
    logger.info(f"PPTX 解析完成: {file_path}, 共 {len(slides_text)} 张幻灯片")
    return result


def parse_md(file_path: str) -> str:
    """
    解析 Markdown 文件
    直接读取原文本, 保留原始格式
    :param file_path: Markdown 文件路径
    :return: 原始文本内容
    """
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()
    logger.info(f"Markdown 解析完成: {file_path}, 共 {len(content)} 字符")
    return content


def parse_txt(file_path: str) -> str:
    """
    解析纯文本文件
    尝试多种编码读取
    :param file_path: TXT 文件路径
    :return: 文本内容
    """
    # 尝试不同编码
    for encoding in ["utf-8", "gbk", "gb2312", "latin-1"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read()
            logger.info(f"TXT 解析完成: {file_path}, 编码={encoding}, 共 {len(content)} 字符")
            return content
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码文件: {file_path}")


def parse_file(file_path: str, file_type: str) -> str:
    """
    统一文件解析入口
    根据文件类型调用对应的解析函数
    :param file_path: 文件的绝对路径
    :param file_type: 文件类型 (pdf/docx/pptx/md/txt)
    :return: 解析后的纯文本内容
    """
    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "pptx": parse_pptx,
        "md": parse_md,
        "txt": parse_txt,
    }
    parser = parsers.get(file_type)
    if not parser:
        raise ValueError(f"不支持的文件类型: {file_type}")
    return parser(file_path)
